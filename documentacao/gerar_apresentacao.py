# -*- coding: utf-8 -*-
"""Gera documentacao/apresentacao_sfm.pptx — apresentação executiva do SFM."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# Paleta (mesma do documento de requisitos)
AZUL_ESCURO = RGBColor(0x0B, 0x2E, 0x4F)
AZUL = RGBColor(0x0B, 0x5C, 0xAB)
AZUL_CLARO = RGBColor(0xEA, 0xF2, 0xFB)
LARANJA = RGBColor(0xE8, 0x75, 0x2C)
CINZA_TEXTO = RGBColor(0x26, 0x31, 0x3D)
CINZA_CLARO = RGBColor(0x6B, 0x7A, 0x89)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
LINHA = RGBColor(0xD9, 0xE2, 0xEC)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color=BRANCO):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def rect(slide, x, y, w, h, color, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h, text, size=18, color=CINZA_TEXTO, bold=False,
            align=PP_ALIGN.LEFT, font="Calibri", anchor=None, line_spacing=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=CINZA_TEXTO, bold_first=False,
            space_after=10, marker="•", font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        p.line_spacing = 1.15
        if isinstance(item, tuple):
            head, rest = item
            r1 = p.add_run()
            r1.text = f"{marker}  {head}"
            r1.font.bold = True
            r1.font.size = Pt(size)
            r1.font.color.rgb = AZUL_ESCURO
            r1.font.name = font
            if rest:
                r2 = p.add_run()
                r2.text = f" — {rest}"
                r2.font.size = Pt(size)
                r2.font.color.rgb = color
                r2.font.name = font
        else:
            r = p.add_run()
            r.text = f"{marker}  {item}"
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.name = font
    return tb


def header(slide, kicker, title, num=None):
    rect(slide, 0, 0, SLIDE_W, Inches(1.15), AZUL_ESCURO)
    rect(slide, 0, Inches(1.15), SLIDE_W, Pt(4), LARANJA)
    textbox(slide, Inches(0.55), Inches(0.13), Inches(9), Inches(0.35), kicker,
            size=13, color=LARANJA, bold=True)
    textbox(slide, Inches(0.55), Inches(0.42), Inches(11.5), Inches(0.65), title,
            size=26, color=BRANCO, bold=True)
    textbox(slide, Inches(0.4), Inches(7.08), Inches(2), Inches(0.35),
            "SFM — Shop Floor Management", size=10, color=CINZA_CLARO)
    if num:
        textbox(slide, Inches(12.4), Inches(7.08), Inches(0.7), Inches(0.35),
                str(num), size=10, color=CINZA_CLARO, align=PP_ALIGN.RIGHT)


def table_slide(slide, x, y, w, h, headers, rows, col_widths=None, font_size=13,
                 header_bg=AZUL, zebra=AZUL_CLARO):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gshape = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    table = gshape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Emu(int(w * cw / total))
    for c, htext in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = htext
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_bg
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(font_size)
                r.font.color.rgb = BRANCO
        cell.margin_top = Pt(4)
        cell.margin_bottom = Pt(4)
        cell.margin_left = Pt(6)
        cell.margin_right = Pt(6)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = zebra if r_idx % 2 == 0 else BRANCO
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for rr in p.runs:
                    rr.font.size = Pt(font_size)
                    rr.font.color.rgb = CINZA_TEXTO
            cell.margin_top = Pt(3)
            cell.margin_bottom = Pt(3)
            cell.margin_left = Pt(6)
            cell.margin_right = Pt(6)
    return table


def stat_card(slide, x, y, w, h, number, label, color=AZUL):
    card = rect(slide, x, y, w, h, AZUL_CLARO)
    rect(slide, x, y, w, Pt(4), color)
    textbox(slide, x + Inches(0.1), y + Inches(0.12), w - Inches(0.2), Inches(0.7),
            number, size=30, color=color, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, x + Inches(0.1), y + h - Inches(0.55), w - Inches(0.2), Inches(0.5),
            label, size=12.5, color=CINZA_TEXTO, align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------------------
# Slide 1 — Capa
# ---------------------------------------------------------------------------
s = add_slide()
set_bg(s, AZUL_ESCURO)
rect(s, 0, Inches(6.6), SLIDE_W, Inches(0.9), AZUL)
rect(s, 0, Inches(6.6), SLIDE_W, Pt(4), LARANJA)
textbox(s, Inches(1), Inches(2.3), Inches(11.3), Inches(1.3), "SFM",
        size=72, color=BRANCO, bold=True, align=PP_ALIGN.CENTER)
textbox(s, Inches(1), Inches(3.55), Inches(11.3), Inches(0.6), "Shop Floor Management",
        size=22, color=LARANJA, bold=True, align=PP_ALIGN.CENTER)
textbox(s, Inches(1), Inches(4.35), Inches(11.3), Inches(0.6),
        "Passagem de turno e painel SQDC — visão geral do projeto", size=17,
        color=RGBColor(0xC9, 0xD8, 0xE8), align=PP_ALIGN.CENTER)
textbox(s, Inches(1), Inches(6.72), Inches(11.3), Inches(0.5),
        "Documento de apresentação  |  Versão 1.0  |  18/09/2026", size=13,
        color=BRANCO, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# Slide 2 — Agenda
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "SFM — VISÃO GERAL", "Agenda", 2)
itens = [
    "O problema e o objetivo do projeto",
    "Escopo: o que está dentro e fora",
    "Como era antes x como é hoje",
    "Módulos e papéis de usuário",
    "Regras de negócio essenciais",
    "Modelo de dados (resumo)",
    "Segurança e limitações conhecidas",
    "Riscos e roadmap",
]
bullets(s, Inches(0.8), Inches(1.6), Inches(11.5), Inches(5.2), itens, size=19, space_after=16)

# ---------------------------------------------------------------------------
# Slide 3 — Problema
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "CONTEXTO", "O problema que o SFM resolve", 3)
bullets(s, Inches(0.8), Inches(1.55), Inches(11.7), Inches(4.7), [
    ("Informação dispersa", "máquina parada e o que já foi tentado ficam registrados em "
     "papel, conversa ou memória do operador, sem histórico entre turnos."),
    ("Indicadores manuais", "\"corretivas atendidas\" e \"quebras graves\" eram calculados "
     "olhando a planilha do SAP linha por linha."),
    ("Quadro físico sem histórico", "o quadro SQDC pendurado no setor não tinha versão "
     "digital nem histórico consultável por dia."),
    ("Sem rastreabilidade", "não havia registro de quem passou e quem recebeu cada máquina, "
     "nem de quanto tempo ela ficou parada de fato."),
], size=17, space_after=18)

# ---------------------------------------------------------------------------
# Slide 4 — Objetivo
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "OBJETIVO", "Objetivo geral do SFM", 4)
box = rect(s, Inches(1.1), Inches(1.9), Inches(11.1), Inches(1.6), AZUL_CLARO)
rect(s, Inches(1.1), Inches(1.9), Pt(6), Inches(1.6), LARANJA)
textbox(s, Inches(1.5), Inches(2.05), Inches(10.3), Inches(1.35),
        "Centralizar e padronizar o controle de passagem de turno das máquinas paradas e o "
        "acompanhamento diário do quadro SQDC por setor, eliminando cálculo manual de "
        "indicadores e dando rastreabilidade completa a cada parada e a cada passagem entre "
        "turnos.", size=18, color=AZUL_ESCURO, bold=True, anchor=MSO_ANCHOR.MIDDLE)
stat_card(s, Inches(1.1), Inches(4.0), Inches(2.65), Inches(1.9), "0", "Servidor externo\n(app 100% local)")
stat_card(s, Inches(3.95), Inches(4.0), Inches(2.65), Inches(1.9), "4", "Setores cobertos\nGasolina, Diesel,\nControle, Biela")
stat_card(s, Inches(6.8), Inches(4.0), Inches(2.65), Inches(1.9), "3", "Turnos monitorados\nManhã, Tarde, Noite")
stat_card(s, Inches(9.65), Inches(4.0), Inches(2.65), Inches(1.9), "S Q D C", "Blocos do quadro\nde indicadores diário")

# ---------------------------------------------------------------------------
# Slide 5 — Escopo
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "ESCOPO DO PRODUTO", "O que está dentro e fora do escopo atual", 5)
textbox(s, Inches(0.8), Inches(1.5), Inches(5.6), Inches(0.4), "Dentro do escopo", size=17,
        bold=True, color=AZUL)
bullets(s, Inches(0.8), Inches(1.95), Inches(5.9), Inches(4.7), [
    "Login local e controle de papel/setor/turno",
    "Passagem e recebimento de turno (ciclo completo)",
    "Importação da planilha do SAP com deduplicação",
    "Quadro SQDC por setor, com D e C automáticos",
    "Relatórios consolidados para gestor/admin",
    "Administração de usuários",
], size=15, space_after=10)
textbox(s, Inches(6.9), Inches(1.5), Inches(5.6), Inches(0.4), "Fora do escopo atual", size=17,
        bold=True, color=LARANJA)
bullets(s, Inches(6.9), Inches(1.95), Inches(5.9), Inches(4.7), [
    "Servidor, nuvem ou autenticação corporativa (SSO)",
    "Aplicativo mobile dedicado",
    "Integração automática/API com o SAP",
    "Fonte automática de dados de Segurança/Qualidade",
    "Criptografia de dados em repouso e auditoria formal",
], size=15, space_after=10, marker="✕", color=CINZA_TEXTO)

# ---------------------------------------------------------------------------
# Slide 6 — Antes x Depois
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "PROCESSO", "Como era antes × como é com o SFM", 6)
rect(s, Inches(0.8), Inches(1.55), Inches(5.7), Inches(5.3), RGBColor(0xF4, 0xF0, 0xEC))
rect(s, Inches(0.8), Inches(1.55), Inches(5.7), Inches(0.6), CINZA_CLARO)
textbox(s, Inches(0.8), Inches(1.6), Inches(5.7), Inches(0.5), "Antes (AS IS)", size=17,
        bold=True, color=BRANCO, align=PP_ALIGN.CENTER)
bullets(s, Inches(1.1), Inches(2.35), Inches(5.1), Inches(4.3), [
    "Comunicação verbal ou anotação entre turnos",
    "Sem registro estruturado de início de parada",
    "Cálculo manual de indicadores olhando o SAP",
    "Quadro SQDC só em papel, sem histórico digital",
], size=14.5, space_after=14)

rect(s, Inches(6.85), Inches(1.55), Inches(5.7), Inches(5.3), AZUL_CLARO)
rect(s, Inches(6.85), Inches(1.55), Inches(5.7), Inches(0.6), AZUL)
textbox(s, Inches(6.85), Inches(1.6), Inches(5.7), Inches(0.5), "Depois (TO BE — implementado)",
        size=17, bold=True, color=BRANCO, align=PP_ALIGN.CENTER)
bullets(s, Inches(7.15), Inches(2.35), Inches(5.1), Inches(4.3), [
    "Sistema redireciona obrigatoriamente para receber/passar turno na janela certa",
    "Cada parada tem início, status e responsável registrados",
    "D e C do SQDC calculados automaticamente pela planilha do SAP",
    "Quadro SQDC digital, por dia, com impressão e histórico",
], size=14.5, space_after=14)

# ---------------------------------------------------------------------------
# Slide 7 — Ciclo de vida da passagem
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "PROCESSO", "Ciclo de vida de uma máquina parada", 7)
labels = ["ABERTA", "RECEBIDA", "FINALIZADA"]
colors_cycle = [LARANJA, AZUL, AZUL_ESCURO]
box_w = Inches(3.0)
gap = Inches(0.9)
start_x = Inches(0.9)
y = Inches(2.6)
for i, (lab, col) in enumerate(zip(labels, colors_cycle)):
    x = start_x + i * (box_w + gap)
    rect(s, x, y, box_w, Inches(1.3), col)
    textbox(s, x, y + Inches(0.42), box_w, Inches(0.5), lab, size=20, bold=True, color=BRANCO,
            align=PP_ALIGN.CENTER)
    if i < 2:
        arrow_x = x + box_w
        tb = textbox(s, arrow_x, y + Inches(0.35), gap, Inches(0.6), "→", size=30,
                      color=CINZA_TEXTO, align=PP_ALIGN.CENTER)
# volta de recebida -> aberta
textbox(s, start_x + box_w + Inches(0.05), Inches(4.05), box_w + gap, Inches(0.5),
        "\"Continuar parada\" volta para ABERTA", size=13, color=CINZA_CLARO,
        align=PP_ALIGN.CENTER)
bullets(s, Inches(0.9), Inches(4.7), Inches(11.5), Inches(2.2), [
    ("Uma única linha por máquina", "não se cria um novo registro a cada troca de turno "
     "(RN-01)."),
    ("Tempo parado", "diferença entre o início real da parada e o momento da finalização — "
     "o \"relógio\" nunca reinicia ao repassar (RN-02)."),
    ("Quebra grave", "toda ordem finalizada com 10h ou mais de parada entra no indicador de "
     "Custo (C) e no Top problemas do dia (RN-03)."),
], size=15, space_after=10)

# ---------------------------------------------------------------------------
# Slide 8 — Módulos
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "ARQUITETURA FUNCIONAL", "Módulos do sistema", 8)
mods = [
    ["Autenticação", "login.html", "Login local por usuário/senha"],
    ["Menu", "menu.html", "Escolha de setor e da ação"],
    ["Passar turno", "passar-turno.html", "Importa SAP e registra paradas novas"],
    ["Receber turno", "receber-turno.html", "Recebe em lote as passagens do setor"],
    ["Quadro SQDC", "quadro-sfm.html", "Folha SQDC do mês, por setor e dia"],
    ["Relatórios", "relatorios.html", "Visão consolidada de todos os setores"],
    ["Administração", "admin.html", "Cadastro de usuários (só admin)"],
]
table_slide(s, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4),
            ["Módulo", "Tela", "Finalidade"], mods,
            col_widths=[2.2, 2.6, 5.2], font_size=14)

# ---------------------------------------------------------------------------
# Slide 9 — Papéis de usuário
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "PERFIS DE ACESSO", "Quem faz o quê no SFM", 9)
perfis = [
    ["Operador", "Passa/recebe turno do próprio setor; se estiver no turno Manhã, marca S/Q "
                 "do quadro do próprio setor.", "Só vê e opera o próprio setor"],
    ["Gestor", "Acompanha o quadro SQDC de qualquer setor na reunião diária; acessa "
               "relatórios consolidados.", "Pode corrigir qualquer dia/setor"],
    ["Administrador", "Gerencia usuários (papel, setor, turno, senha); mesmos acessos do "
                       "gestor.", "Único perfil com Administração"],
]
table_slide(s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(3.4),
            ["Perfil", "Pode fazer", "Restrição principal"], perfis,
            col_widths=[1.8, 5.8, 3.5], font_size=14.5)
textbox(s, Inches(0.8), Inches(5.35), Inches(11.7), Inches(1.4),
        "Reunião diária de SFM: segunda a sexta, às 07:40, conduzida pelo gestor, que "
        "acompanha o quadro SQDC de cada setor sem precisar editar nada.", size=15,
        color=AZUL_ESCURO, bold=True)

# ---------------------------------------------------------------------------
# Slide 10 — Regras de negócio chave
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "REGRAS DE NEGÓCIO", "Regras essenciais do sistema", 10)
regras = [
    ["RN-01", "Uma parada = um registro só, que muda de status ao longo do tempo"],
    ["RN-04", "Atendida/Não atendida sempre calculada pelo Status sistema do SAP"],
    ["RN-05", "Importar planilha nunca duplica nota — atualiza a que já existe"],
    ["RN-06", "Qualquer operador do turno Manhã pode marcar S/Q do próprio setor"],
    ["RN-07", "S/Q só pode ser marcado nos dias que a janela da SFM cobre"],
    ["RN-08", "Janelas obrigatórias de 40 min redirecionam o operador automaticamente"],
]
table_slide(s, Inches(0.8), Inches(1.55), Inches(11.7), Inches(5.2),
            ["ID", "Regra"], regras, col_widths=[1.2, 9.5], font_size=15)

# ---------------------------------------------------------------------------
# Slide 11 — Modelo de dados
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "DADOS", "Modelo de dados (resumo)", 11)
textbox(s, Inches(0.8), Inches(1.45), Inches(11.7), Inches(0.6),
        "Sem banco relacional: cada arquivo .json é uma coleção de registros, com um "
        "espelho .js regenerado a cada gravação.", size=15, color=CINZA_TEXTO)
dados = [
    ["bd/dados.json", "usuarios, passagensTurno,\neventosPassagem, confirmacoesTurno",
     "Usuários, estado das paradas, histórico de passagens"],
    ["bd/notas.json", "lista de notas do SAP",
     "Uma linha por número de Nota, atualizada a cada importação"],
    ["bd/quadro.json", "registros por setor + dia",
     "Marcações manuais de Segurança e Qualidade"],
]
table_slide(s, Inches(0.8), Inches(2.2), Inches(11.7), Inches(2.6),
            ["Arquivo", "Coleções", "Conteúdo"], dados,
            col_widths=[2.6, 3.6, 5.5], font_size=14)
textbox(s, Inches(0.8), Inches(5.1), Inches(11.7), Inches(1.6),
        "Cada entidade principal tem identificador interno automático (ex.: passagensTurno "
        "usa \"pt-\" + uuid), e campos de auditoria mínima como registradoPor/recebidoPor e "
        "datas — mas não há criptografia nem trilha de auditoria de leitura.", size=14,
        color=CINZA_CLARO)

# ---------------------------------------------------------------------------
# Slide 12 — Segurança e limitações
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "SEGURANÇA", "Segurança e limitações conhecidas", 12)
textbox(s, Inches(0.8), Inches(1.5), Inches(5.6), Inches(0.4), "Existe hoje", size=17,
        bold=True, color=AZUL)
bullets(s, Inches(0.8), Inches(1.95), Inches(5.9), Inches(4.7), [
    "Senha com hash PBKDF2 + salt",
    "Acesso por papel e por setor",
    "Registro de quem passou/recebeu cada ordem",
    "Aviso de conflito de gravação concorrente",
], size=15, space_after=12)
textbox(s, Inches(6.9), Inches(1.5), Inches(5.6), Inches(0.4), "Limitação aceita", size=17,
        bold=True, color=LARANJA)
bullets(s, Inches(6.9), Inches(1.95), Inches(5.9), Inches(4.7), [
    "Não é segurança de verdade — dados legíveis por quem acessa a pasta de rede",
    "Sem criptografia em repouso",
    "Sem trilha de auditoria de leitura",
    "Sem autenticação multifator ou expiração de sessão",
], size=15, space_after=12, marker="!", color=CINZA_TEXTO)
box = rect(s, Inches(0.8), Inches(6.35), Inches(11.7), Inches(0.75), AZUL_CLARO)
textbox(s, Inches(1.0), Inches(6.42), Inches(11.3), Inches(0.6),
        "Recomendação: não guardar dados sensíveis no sistema e trocar as senhas de teste "
        "antes do uso real.", size=13.5, color=AZUL_ESCURO, bold=True,
        anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------------------
# Slide 13 — Riscos
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "RISCOS", "Riscos e pontos de atenção", 13)
riscos = [
    ["Sem servidor/backup central", "Alto", "Backup regular da pasta de rede pela TI"],
    ["Conflito de gravação simultânea", "Médio", "Detecção já existe; resolução é manual"],
    ["Coluna de data ausente na planilha SAP", "Alto", "Diagnóstico de cabeçalhos na importação"],
    ["Relógio do computador incorreto", "Médio", "Sincronização automática de hora (NTP)"],
]
table_slide(s, Inches(0.8), Inches(1.55), Inches(11.7), Inches(3.5),
            ["Risco", "Impacto", "Mitigação"], riscos, col_widths=[4.5, 1.8, 5.4],
            font_size=14)

# ---------------------------------------------------------------------------
# Slide 14 — Roadmap
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s)
header(s, "PRÓXIMOS PASSOS", "Roadmap sugerido", 14)
fases = ["Consolidação\n(atual)", "Qualidade\nde dados", "Fonte automática\nde S/Q",
         "Robustez\noperacional", "Limpeza\ntécnica"]
box_w = Inches(2.15)
gap = Inches(0.25)
start_x = Inches(0.7)
y = Inches(1.8)
for i, f in enumerate(fases):
    x = start_x + i * (box_w + gap)
    col = AZUL if i % 2 == 0 else AZUL_ESCURO
    rect(s, x, y, box_w, Inches(1.1), col)
    textbox(s, x + Inches(0.05), y + Inches(0.12), box_w - Inches(0.1), Inches(0.9), f,
            size=13.5, bold=True, color=BRANCO, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
    if i < len(fases) - 1:
        textbox(s, x + box_w, y + Inches(0.3), gap, Inches(0.5), "→", size=18,
                color=CINZA_CLARO, align=PP_ALIGN.CENTER)
bullets(s, Inches(0.8), Inches(3.4), Inches(11.7), Inches(3.5), [
    "Documentar e treinar novos usuários; trocar senhas de teste antes de expandir o uso",
    "Ampliar diagnóstico de importação da planilha do SAP",
    "Avaliar fonte automática para Segurança e Qualidade (S/Q), hoje 100% manual",
    "Rotina de backup automática da pasta bd/ e melhorar tratamento de conflito de gravação",
    "Revisar campos legados (responsavelSfm, eficiencia) para reduzir dívida técnica",
], size=15.5, space_after=12)

# ---------------------------------------------------------------------------
# Slide 15 — Encerramento
# ---------------------------------------------------------------------------
s = add_slide(); set_bg(s, AZUL_ESCURO)
rect(s, 0, Inches(3.4), SLIDE_W, Pt(4), LARANJA)
textbox(s, Inches(1), Inches(2.6), Inches(11.3), Inches(0.7), "Obrigado.", size=40,
        bold=True, color=BRANCO, align=PP_ALIGN.CENTER)
textbox(s, Inches(1), Inches(3.6), Inches(11.3), Inches(0.6),
        "Dúvidas e sugestões: equipe do projeto SFM", size=16,
        color=RGBColor(0xC9, 0xD8, 0xE8), align=PP_ALIGN.CENTER)

prs.save("apresentacao_sfm.pptx")
print("OK: apresentacao_sfm.pptx gerado —", len(prs.slides._sldIdLst), "slides")
